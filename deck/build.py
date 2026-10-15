"""Build the deck: 13 slides, one design system, restrained animation."""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

from . import theme as T
from . import data as D
from . import qr
from .helpers import (rgb, blank, fill_bg, rect, line, textbox, para, runs,
                      add_title, add_footer, icon, bar_chart, line_chart,
                      animate)

TOTAL = 13
REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(REPO, "out")            # scratch: icons, qr, previews
DELIVERABLE = os.path.join(REPO, "future-of-work.pptx")


def _prs():
    prs = Presentation()
    prs.slide_width = Inches(T.SLIDE_W)
    prs.slide_height = Inches(T.SLIDE_H)
    return prs


# --- slide 1: cover ---------------------------------------------------------
def slide_cover(prs):
    s = blank(prs)
    fill_bg(s, T.NAVY)
    # thin accent stripe
    rect(s, 0, 0, T.SLIDE_W, 0.14, fill=T.GREEN)
    # brand mark (vector)
    icon(s, "brain", T.WHITE, T.SLIDE_W - T.MARGIN - 1.05, T.MARGIN - 0.05, 1.05)
    tb, tf = textbox(s, T.MARGIN, 2.15, 9.6, 1.9)
    para(tf, D.TITLE, size=T.SZ_COVER_TITLE, color=T.WHITE, bold=True,
         first=True, line_spacing=1.02)
    line(s, T.MARGIN, 4.15, T.MARGIN + 2.4, 4.15, color=T.GREEN, width=4.0)
    tb2, tf2 = textbox(s, T.MARGIN, 4.4, 9.6, 1.0)
    para(tf2, D.SUBTITLE, size=T.SZ_COVER_SUB, color="DCE6F2", first=True,
         line_spacing=1.1)
    tb3, tf3 = textbox(s, T.MARGIN, 6.05, 9.6, 1.0)
    p = para(tf3, "", size=T.SZ_SMALL, first=True)
    runs(p, [(D.AUTHOR + "   ", {"color": T.WHITE, "bold": True}),
             (D.ORG + "   ·   " + D.DATE, {"color": "9FB3C8"})],
         size=T.SZ_SMALL)
    return s


# --- slide 2: contents ------------------------------------------------------
def slide_toc(prs):
    s = blank(prs)
    fill_bg(s)
    add_title(s, "Содержание")
    top = 1.95
    gap = 0.86
    for i, (name, hint) in enumerate(D.TOC):
        y = top + i * gap
        marker = rect(s, T.MARGIN, y, 0.62, 0.62, fill=T.NAVY,
                      shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.22)
        mtb, mtf = textbox(s, T.MARGIN, y, 0.62, 0.62,
                           anchor=MSO_ANCHOR.MIDDLE)
        para(mtf, str(i + 1), size=24, color=T.WHITE, bold=True,
             align=PP_ALIGN.CENTER, first=True)
        tb, tf = textbox(s, T.MARGIN + 0.95, y, 7.6, 0.62,
                         anchor=MSO_ANCHOR.MIDDLE)
        p = para(tf, "", first=True)
        runs(p, [(name, {"bold": True, "color": T.NAVY, "size": T.SZ_BODY}),
                 ("  —  " + hint, {"color": T.GREY, "size": T.SZ_SMALL})])
        if i < len(D.TOC) - 1:
            line(s, T.MARGIN + 0.95, y + 0.74, T.SLIDE_W - T.MARGIN, y + 0.74,
                 color=T.GREY_LIGHT, width=0.75)
        animate(s, marker.shape_id, "wipe_right", delay=0,
                trigger="click" if i == 0 else "after")
    add_footer(s, 2, TOTAL)
    return s


# --- slide 3: scale of the problem ------------------------------------------
def slide_intro(prs):
    s = blank(prs)
    fill_bg(s)
    add_title(s, "Автоматизация неизбежна?")
    tb, tf = textbox(s, T.MARGIN, 1.72, T.CONTENT_W, 0.5)
    para(tf, "До трети рабочего времени и половины рабочих мест "
             "уже находятся в зоне автоматизации.", size=T.SZ_SMALL,
         color=T.GREY, first=True)
    cw = (T.CONTENT_W - 0.5) / 2.0
    for i, st in enumerate(D.STATS):
        x = T.MARGIN + i * (cw + 0.5)
        color = T.NAVY if i == 0 else T.ORANGE
        tint = T.NAVY_TINT if i == 0 else T.ORANGE_TINT
        card = rect(s, x, 2.35, cw, 3.35, fill=T.WHITE, line=T.GREY_LIGHT,
                    line_w=1.0)
        rect(s, x, 2.35, cw, 0.12, fill=color)
        vtb, vtf = textbox(s, x + 0.4, 2.7, cw - 0.8, 1.7,
                           anchor=MSO_ANCHOR.MIDDLE)
        para(vtf, st["value"], size=T.SZ_BIG, color=color, bold=True,
             align=PP_ALIGN.CENTER, first=True)
        ltb, ltf = textbox(s, x + 0.4, 4.45, cw - 0.8, 0.5,
                           anchor=MSO_ANCHOR.MIDDLE)
        para(ltf, st["label"], size=T.SZ_SUBTITLE - 4, color=T.TEXT,
             bold=True, align=PP_ALIGN.CENTER, first=True)
        dtb, dtf = textbox(s, x + 0.4, 5.05, cw - 0.8, 0.7,
                           anchor=MSO_ANCHOR.TOP)
        para(dtf, st["text"], size=T.SZ_SMALL, color=T.TEXT,
             align=PP_ALIGN.CENTER, first=True, line_spacing=1.05)
        stb, stf = textbox(s, x + 0.4, 5.72, cw - 0.8, 0.4)
        para(stf, st["source"], size=T.SZ_SMALL, color=T.GREY,
             align=PP_ALIGN.CENTER, first=True)
        animate(s, card.shape_id, "zoom", delay=0,
                trigger="click" if i == 0 else "after")
    add_footer(s, 3, TOTAL)
    return s


# --- slide 4: methodology ---------------------------------------------------
def slide_factors(prs):
    s = blank(prs)
    fill_bg(s)
    add_title(s, "Критерии уязвимости профессий")
    tb, tf = textbox(s, T.MARGIN, 1.7, T.CONTENT_W, 0.5)
    para(tf, "Пять факторов, по которым оценивают, насколько профессия "
             "подвержена автоматизации.", size=T.SZ_SMALL, color=T.GREY,
         first=True)
    cols = 5
    gap = 0.28
    cw = (T.CONTENT_W - gap * (cols - 1)) / cols
    top = 2.55
    for i, (ic, title, desc) in enumerate(D.FACTORS):
        x = T.MARGIN + i * (cw + gap)
        card = rect(s, x, top, cw, 3.5, fill=T.WHITE, line=T.GREY_LIGHT,
                    line_w=1.0)
        rect(s, x, top, cw, 0.09, fill=T.NAVY)
        icon(s, ic, T.NAVY, x + (cw - 1.05) / 2, top + 0.35, 1.05)
        ntb, ntf = textbox(s, x + 0.12, top + 1.6, cw - 0.24, 0.9,
                           anchor=MSO_ANCHOR.TOP)
        para(ntf, title, size=T.SZ_SMALL + 1, color=T.NAVY, bold=True,
             align=PP_ALIGN.CENTER, first=True, line_spacing=1.0)
        dtb, dtf = textbox(s, x + 0.12, top + 2.5, cw - 0.24, 0.9)
        para(dtf, desc, size=T.SZ_SMALL, color=T.GREY,
             align=PP_ALIGN.CENTER, first=True, line_spacing=1.05)
        animate(s, card.shape_id, "fade", delay=0,
                trigger="click" if i == 0 else "after")
    add_footer(s, 4, TOTAL)
    return s


# --- slide 5: high-risk jobs ------------------------------------------------
def slide_high(prs):
    s = blank(prs)
    fill_bg(s)
    add_title(s, "Под ударом: высокий риск автоматизации", accent=T.RED)
    top = 1.85
    gap = 1.02
    for i, (ic, name, desc) in enumerate(D.HIGH_RISK):
        y = top + i * gap
        card = rect(s, T.MARGIN, y, T.CONTENT_W, 0.88, fill=T.WHITE,
                    line=T.GREY_LIGHT, line_w=1.0)
        rect(s, T.MARGIN, y, 0.14, 0.88, fill=T.RED)
        icon(s, ic, T.RED, T.MARGIN + 0.34, y + 0.14, 0.6)
        tb, tf = textbox(s, T.MARGIN + 1.2, y, 5.6, 0.88,
                         anchor=MSO_ANCHOR.MIDDLE)
        para(tf, name, size=T.SZ_BODY, color=T.NAVY, bold=True, first=True)
        dtb, dtf = textbox(s, T.MARGIN + 7.0, y, T.CONTENT_W - 7.2, 0.88,
                           anchor=MSO_ANCHOR.MIDDLE)
        para(dtf, desc, size=T.SZ_SMALL, color=T.GREY, first=True,
             align=PP_ALIGN.RIGHT)
        animate(s, card.shape_id, "wipe_left", delay=0,
                trigger="click" if i == 0 else "after")
    add_footer(s, 5, TOTAL, note="Акцент — риск автоматизации")
    return s


# --- slide 6: bar chart -----------------------------------------------------
def slide_bar(prs):
    s = blank(prs)
    fill_bg(s)
    add_title(s, "Насколько автоматизируемы задачи?")
    tb, tf = textbox(s, T.MARGIN, 1.7, T.CONTENT_W, 0.5)
    para(tf, "Вероятность автоматизации профессий, % задач", size=T.SZ_SMALL,
         color=T.GREY, first=True)
    chart = bar_chart(s, T.MARGIN + 0.1, 2.2, T.CONTENT_W - 0.2, 4.25,
                      D.CHART_CATS, "Доля автоматизируемых задач, %",
                      D.CHART_VALUES,
                      colors=[T.RED, T.RED, T.ORANGE, T.RED, T.ORANGE])
    tb2, tf2 = textbox(s, T.MARGIN, 6.55, T.CONTENT_W, 0.4)
    para(tf2, D.CHART_SOURCE, size=T.SZ_SMALL, color=T.GREY, first=True)
    add_footer(s, 6, TOTAL)
    return s


# --- slide 7: low-risk jobs -------------------------------------------------
def slide_low(prs):
    s = blank(prs)
    fill_bg(s)
    add_title(s, "Вне зоны риска: устойчивые профессии", accent=T.GREEN)
    top = 1.85
    gap = 1.02
    for i, (ic, name, desc) in enumerate(D.LOW_RISK):
        y = top + i * gap
        card = rect(s, T.MARGIN, y, T.CONTENT_W, 0.88, fill=T.WHITE,
                    line=T.GREY_LIGHT, line_w=1.0)
        rect(s, T.MARGIN, y, 0.14, 0.88, fill=T.GREEN)
        icon(s, ic, T.GREEN, T.MARGIN + 0.34, y + 0.14, 0.6)
        tb, tf = textbox(s, T.MARGIN + 1.2, y, 6.0, 0.88,
                         anchor=MSO_ANCHOR.MIDDLE)
        para(tf, name, size=T.SZ_BODY, color=T.NAVY, bold=True, first=True)
        dtb, dtf = textbox(s, T.MARGIN + 7.4, y, T.CONTENT_W - 7.6, 0.88,
                           anchor=MSO_ANCHOR.MIDDLE)
        para(dtf, desc, size=T.SZ_SMALL, color=T.GREY, first=True,
             align=PP_ALIGN.RIGHT)
        animate(s, card.shape_id, "fade", delay=0,
                trigger="click" if i == 0 else "after")
    add_footer(s, 7, TOTAL, note="Акцент — устойчивость к автоматизации")
    return s


# --- slide 8: comparison table ----------------------------------------------
def slide_table(prs):
    s = blank(prs)
    fill_bg(s)
    add_title(s, "Риски и возможности")
    rows = len(D.TABLE_ROWS) + 1
    cols = 3
    x, y = T.MARGIN, 1.95
    w, h = T.CONTENT_W * 0.9, 4.3
    gt = s.shapes.add_table(rows, cols, Inches(x), Inches(y),
                            Inches(w), Inches(h))
    tbl = gt.table
    tbl.columns[0].width = Inches(w * 0.24)
    tbl.columns[1].width = Inches(w * 0.24)
    tbl.columns[2].width = Inches(w * 0.52)
    # header
    for c, text in enumerate(D.TABLE_HEAD):
        cell = tbl.cell(0, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = rgb(T.NAVY)
        cell.margin_left = cell.margin_right = Inches(0.12)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER
        r = p.add_run()
        r.text = text
        r.font.size = Pt(T.SZ_SMALL)
        r.font.bold = True
        r.font.name = T.FONT
        r.font.color.rgb = rgb(T.WHITE)
    for ri, (prof, risk, opp) in enumerate(D.TABLE_ROWS, start=1):
        zebra = T.WHITE if ri % 2 else T.GREY_LIGHT
        risk_color = T.RED if risk == "Высокий" else T.GREEN
        for ci, text in enumerate((prof, risk, opp)):
            cell = tbl.cell(ri, ci)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(zebra)
            cell.margin_left = cell.margin_right = Inches(0.12)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]
            p.alignment = (PP_ALIGN.LEFT if ci == 0
                           else PP_ALIGN.CENTER if ci == 1
                           else PP_ALIGN.LEFT)
            r = p.add_run()
            r.text = text
            r.font.size = Pt(T.SZ_SMALL if ci != 1 else T.SZ_SMALL)
            r.font.bold = ci in (0, 1)
            r.font.name = T.FONT
            r.font.color.rgb = rgb(T.NAVY if ci == 0
                                   else risk_color if ci == 1 else T.TEXT)
    tb, tf = textbox(s, T.MARGIN, 6.4, T.CONTENT_W, 0.4)
    para(tf, "Одна и та же профессия может терять рутинные задачи и "
             "получать новые роли.", size=T.SZ_SMALL, color=T.GREY,
         first=True)
    # animate the table as one object (row-by-row is not exposed by the API)
    animate(s, gt.shape_id, "fade", delay=0, trigger="click")
    add_footer(s, 8, TOTAL)
    return s


# --- slide 9: new professions -----------------------------------------------
def slide_new_jobs(prs):
    s = blank(prs)
    fill_bg(s)
    add_title(s, "Кем станут люди?")
    tb, tf = textbox(s, T.MARGIN, 1.7, T.CONTENT_W, 0.5)
    para(tf, "Вместе с автоматизацией появляются новые профессии "
             "на стыке человека и ИИ.", size=T.SZ_SMALL, color=T.GREY,
         first=True)
    # "human + AI" scheme on the left
    scheme_x = T.MARGIN
    scheme_y = 2.6
    icon(s, "people", T.NAVY, scheme_x, scheme_y + 0.55, 1.15)
    ptb, ptf = textbox(s, scheme_x + 1.2, scheme_y + 1.35, 1.6, 0.4)
    para(ptf, "Человек", size=T.SZ_SMALL, color=T.NAVY, bold=True,
         align=PP_ALIGN.CENTER, first=True)
    # plus
    ptb, ptf = textbox(s, scheme_x + 1.35, scheme_y + 0.85, 0.9, 0.6,
                       anchor=MSO_ANCHOR.MIDDLE)
    para(ptf, "+", size=40, color=T.GREEN, bold=True, align=PP_ALIGN.CENTER,
         first=True)
    icon(s, "robot", T.GREEN, scheme_x + 2.35, scheme_y + 0.55, 1.15)
    ptb, ptf = textbox(s, scheme_x + 2.2, scheme_y + 1.35, 1.5, 0.4)
    para(ptf, "ИИ", size=T.SZ_SMALL, color=T.GREEN, bold=True,
         align=PP_ALIGN.CENTER, first=True)
    ptb, ptf = textbox(s, scheme_x, scheme_y + 2.0, 3.6, 0.6)
    para(ptf, "Новые роли", size=T.SZ_SUBTITLE - 4, color=T.NAVY, bold=True,
         align=PP_ALIGN.CENTER, first=True)
    # jobs on the right
    list_x = T.MARGIN + 4.15
    list_w = T.SLIDE_W - T.MARGIN - list_x
    jy = 2.05
    jgap = 1.02
    for i, (ic, name, desc) in enumerate(D.NEW_JOBS):
        y = jy + i * jgap
        icon(s, ic, T.NAVY, list_x, y + 0.06, 0.66)
        tb, tf = textbox(s, list_x + 0.9, y, list_w - 0.9, 0.78,
                         anchor=MSO_ANCHOR.MIDDLE)
        para(tf, name, size=T.SZ_BODY - 1, color=T.NAVY, bold=True, first=True)
        para(tf, desc, size=T.SZ_SMALL, color=T.GREY)
    add_footer(s, 9, TOTAL)
    return s


# --- slide 10: outlook ------------------------------------------------------
def slide_outlook(prs):
    s = blank(prs)
    fill_bg(s)
    add_title(s, "Что нас ждёт? Прогноз до 2030 года")
    tb, tf = textbox(s, T.MARGIN, 1.7, T.CONTENT_W, 0.5)
    para(tf, "Спрос на профессии, 2024 = 100", size=T.SZ_SMALL, color=T.GREY,
         first=True)
    line_chart(s, T.MARGIN + 0.1, 2.15, T.CONTENT_W - 0.2, 4.2,
               D.OUTLOOK_CATS, D.OUTLOOK_SERIES)
    tb2, tf2 = textbox(s, T.MARGIN, 6.5, T.CONTENT_W, 0.4)
    para(tf2, D.OUTLOOK_SOURCE, size=T.SZ_SMALL, color=T.GREY, first=True)
    add_footer(s, 10, TOTAL)
    return s


# --- slide 11: skills -------------------------------------------------------
def slide_skills(prs):
    s = blank(prs)
    fill_bg(s)
    add_title(s, "Навыки будущего")
    tb, tf = textbox(s, T.MARGIN, 1.7, T.CONTENT_W, 0.5)
    para(tf, "Четыре опоры адаптации к автоматизированной экономике.",
         size=T.SZ_SMALL, color=T.GREY, first=True)
    # 2x2 grid
    gw = (T.CONTENT_W - 0.5) / 2.0
    gh = 1.9
    top = 2.5
    for i, (ic, name, desc) in enumerate(D.SKILLS):
        col = i % 2
        row = i // 2
        x = T.MARGIN + col * (gw + 0.5)
        y = top + row * (gh + 0.45)
        card = rect(s, x, y, gw, gh, fill=T.WHITE, line=T.GREY_LIGHT,
                    line_w=1.0)
        rect(s, x, y, 0.14, gh, fill=T.GREEN)
        icon(s, ic, T.GREEN, x + 0.4, y + 0.42, 1.0)
        tb, tf = textbox(s, x + 1.65, y + 0.28, gw - 1.9, gh - 0.5,
                         anchor=MSO_ANCHOR.MIDDLE)
        para(tf, name, size=T.SZ_BODY, color=T.NAVY, bold=True, first=True)
        para(tf, desc, size=T.SZ_SMALL, color=T.GREY)
        animate(s, card.shape_id, "fade", delay=0,
                trigger="click" if i == 0 else "after")
    add_footer(s, 11, TOTAL)
    return s


# --- slide 12: conclusions --------------------------------------------------
def slide_conclusions(prs):
    s = blank(prs)
    fill_bg(s)
    add_title(s, "Главное")
    top = 2.1
    gap = 1.28
    for i, text in enumerate(D.CONCLUSIONS):
        y = top + i * gap
        rect(s, T.MARGIN, y, 0.12, 1.0, fill=T.GREEN)
        tb, tf = textbox(s, T.MARGIN + 0.55, y, T.CONTENT_W - 0.55, 1.0,
                         anchor=MSO_ANCHOR.MIDDLE)
        p = para(tf, "", first=True)
        runs(p, [(f"{i + 1}.  ", {"bold": True, "color": T.GREEN,
                                  "size": T.SZ_SUBTITLE}),
                 (text, {"color": T.NAVY, "bold": True,
                         "size": T.SZ_SUBTITLE})])
    add_footer(s, 12, TOTAL)
    return s


# --- slide 13: closing ------------------------------------------------------
def slide_closing(prs):
    s = blank(prs)
    fill_bg(s, T.NAVY)
    rect(s, 0, 0, T.SLIDE_W, 0.14, fill=T.GREEN)
    tb, tf = textbox(s, T.MARGIN, 1.35, 8.2, 1.2)
    para(tf, "Спасибо за внимание!", size=T.SZ_COVER_TITLE, color=T.WHITE,
         bold=True, first=True)
    line(s, T.MARGIN, 2.85, T.MARGIN + 2.4, 2.85, color=T.GREEN, width=4.0)
    # sources
    tb2, tf2 = textbox(s, T.MARGIN, 3.1, 8.2, 2.6)
    para(tf2, "Источники", size=T.SZ_SUBTITLE - 6, color=T.GREEN, bold=True,
         first=True, space_after=6)
    for src in D.SOURCES:
        para(tf2, "•  " + src, size=T.SZ_SMALL, color="DCE6F2",
             space_after=4, line_spacing=1.0)
    # contacts + QR
    qr_path = qr.render(D.CONTACT_QR, dark=T.NAVY)
    s.shapes.add_picture(qr_path, Inches(T.SLIDE_W - T.MARGIN - 2.35),
                         Inches(3.0), Inches(2.35), Inches(2.35))
    ctb, ctf = textbox(s, T.SLIDE_W - T.MARGIN - 2.35, 5.45, 2.35, 0.5)
    para(ctf, "Проект на GitHub", size=T.SZ_SMALL, color="DCE6F2",
         align=PP_ALIGN.CENTER, first=True)
    tb3, tf3 = textbox(s, T.MARGIN, 6.4, 9.0, 0.5)
    para(tf3, D.AUTHOR + "   ·   " + D.ORG, size=T.SZ_SMALL,
         color="9FB3C8", first=True)
    return s


def build(path=None):
    prs = _prs()
    slide_cover(prs)
    slide_toc(prs)
    slide_intro(prs)
    slide_factors(prs)
    slide_high(prs)
    slide_bar(prs)
    slide_low(prs)
    slide_table(prs)
    slide_new_jobs(prs)
    slide_outlook(prs)
    slide_skills(prs)
    slide_conclusions(prs)
    slide_closing(prs)
    os.makedirs(OUT, exist_ok=True)
    path = path or DELIVERABLE
    prs.save(path)
    return path


if __name__ == "__main__":
    print(build())