"""Low-level building blocks: shapes, text, rules, charts, animation."""
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_TICK_MARK
from pptx.chart.data import CategoryChartData
from pptx.oxml.ns import qn

from . import theme as T
from . import icons


def rgb(hexstr):
    return RGBColor.from_string(hexstr)


def icon(slide, name, color, x, y, size):
    """Place a vector-style icon (PNG, transparent) at (x, y) inches."""
    path = icons.render(name, color)
    return slide.shapes.add_picture(path, Inches(x), Inches(y),
                                    Inches(size), Inches(size))


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


# --- background -------------------------------------------------------------
def fill_bg(slide, color=T.BG_LIGHT):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0,
                                Inches(T.SLIDE_W), Inches(T.SLIDE_H))
    bg.fill.solid()
    bg.fill.fore_color.rgb = rgb(color)
    bg.line.fill.background()
    bg.shadow.inherit = False
    return bg


def rect(slide, x, y, w, h, fill=None, line=None, line_w=1.0,
         shape=MSO_SHAPE.RECTANGLE, radius=None):
    sh = slide.shapes.add_shape(shape, Inches(x), Inches(y),
                                Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(fill)
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = rgb(line)
        sh.line.width = Pt(line_w)
    sh.shadow.inherit = False
    if radius is not None:
        try:
            sh.adjustments[0] = radius
        except (IndexError, ValueError):
            pass
    return sh


def line(slide, x1, y1, x2, y2, color=T.GREY_LIGHT, width=1.0, dash=None):
    ln = slide.shapes.add_connector(1, Inches(x1), Inches(y1),
                                    Inches(x2), Inches(y2))
    ln.line.color.rgb = rgb(color)
    ln.line.width = Pt(width)
    if dash:
        ln.line.dash_style = dash
    ln.shadow.inherit = False
    return ln


# --- text -------------------------------------------------------------------
def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP, wrap=True):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    return tb, tf


def para(tf, text="", size=T.SZ_BODY, color=T.TEXT, bold=False, italic=False,
         align=PP_ALIGN.LEFT, space_before=0, space_after=0, first=False,
         line_spacing=1.0, font=T.FONT):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    p.line_spacing = line_spacing
    if text != "":
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.name = font
        r.font.color.rgb = rgb(color)
    return p


def runs(p, parts, size=T.SZ_BODY, color=T.TEXT, font=T.FONT):
    """parts: list of (text, {bold, italic, color, size})."""
    for text, opt in parts:
        r = p.add_run()
        r.text = text
        r.font.size = Pt(opt.get("size", size))
        r.font.bold = opt.get("bold", False)
        r.font.italic = opt.get("italic", False)
        r.font.name = font
        r.font.color.rgb = rgb(opt.get("color", color))
    return p


# --- slide chrome -----------------------------------------------------------
def add_title(slide, text, accent=T.NAVY):
    tb, tf = textbox(slide, T.MARGIN, 0.55, T.CONTENT_W, 0.9,
                     anchor=MSO_ANCHOR.MIDDLE)
    para(tf, text, size=T.SZ_TITLE, color=accent, bold=True, first=True)
    line(slide, T.MARGIN, 1.42, T.MARGIN + 2.1, 1.42, color=accent, width=3.0)
    return tb


def add_footer(slide, page, total=13, note=None):
    tb, tf = textbox(slide, T.MARGIN, T.FOOTER_Y, T.CONTENT_W, 0.3)
    parts = []
    if note:
        parts.append((note, {"color": T.GREY, "size": T.SZ_SMALL}))
    p = para(tf, "", size=T.SZ_SMALL, first=True)
    if note:
        runs(p, [(note, {"color": T.GREY, "size": T.SZ_SMALL})])
    tb2, tf2 = textbox(slide, T.SLIDE_W - T.MARGIN - 1.0, T.FOOTER_Y,
                       1.0, 0.3)
    para(tf2, f"{page} / {total}", size=T.SZ_SMALL, color=T.GREY,
         align=PP_ALIGN.RIGHT, first=True)


# --- bullets ----------------------------------------------------------------
def bullet_row(slide, x, y, w, marker_color, marker_text, title, desc,
               msize=30, tsize=T.SZ_BODY, dsize=T.SZ_SMALL, mshape=None):
    """Coloured marker (icon/glyph) + bold title + muted description."""
    box = 0.62
    if mshape is not None:
        rect(slide, x, y, box, box, fill=marker_color,
             shape=mshape, radius=0.22)
        mtb, mtf = textbox(slide, x, y, box, box, anchor=MSO_ANCHOR.MIDDLE)
        para(mtf, marker_text, size=msize - 4, color=T.WHITE, bold=True,
             align=PP_ALIGN.CENTER, first=True)
    else:
        rect(slide, x, y, box, box, fill=marker_color,
             shape=MSO_SHAPE.OVAL)
        mtb, mtf = textbox(slide, x, y, box, box, anchor=MSO_ANCHOR.MIDDLE)
        para(mtf, marker_text, size=msize - 6, color=T.WHITE, bold=True,
             align=PP_ALIGN.CENTER, first=True)
    tb, tf = textbox(slide, x + box + 0.28, y - 0.02, w - box - 0.28, box + 0.1,
                     anchor=MSO_ANCHOR.MIDDLE)
    para(tf, title, size=tsize, color=T.NAVY, bold=True, first=True)
    if desc:
        para(tf, desc, size=dsize, color=T.GREY)


# --- charts -----------------------------------------------------------------
def _style_chart(chart, font_size=T.SZ_SMALL):
    chart.font.size = Pt(font_size)
    chart.font.name = T.FONT
    chart.font.color.rgb = rgb(T.TEXT)


def bar_chart(slide, x, y, w, h, cats, series_name, values, colors=None,
              vmax=100):
    cd = CategoryChartData()
    cd.categories = cats
    cd.add_series(series_name, values)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x),
                                Inches(y), Inches(w), Inches(h), cd)
    chart = gf.chart
    chart.has_legend = False
    chart.has_title = False
    _style_chart(chart)
    plot = chart.plots[0]
    plot.gap_width = 55
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.font.size = Pt(T.SZ_SMALL)
    dl.font.bold = True
    dl.font.name = T.FONT
    dl.font.color.rgb = rgb(T.TEXT)
    va = chart.value_axis
    va.maximum_scale = vmax
    va.minimum_scale = 0
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = rgb(T.GREY_LIGHT)
    va.major_gridlines.format.line.width = Pt(0.75)
    va.format.line.fill.background()
    va.tick_labels.font.size = Pt(T.SZ_SMALL)
    ca = chart.category_axis
    ca.has_major_gridlines = False
    ca.format.line.color.rgb = rgb(T.GREY)
    ca.tick_labels.font.size = Pt(T.SZ_SMALL)
    if colors:
        for i, c in enumerate(colors):
            pt = plot.series[0].points[i]
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = rgb(c)
    else:
        plot.series[0].format.fill.solid()
        plot.series[0].format.fill.fore_color.rgb = rgb(T.NAVY)
    return chart


def line_chart(slide, x, y, w, h, cats, series):
    """series: list of (name, values, color, dashed)."""
    cd = CategoryChartData()
    cd.categories = cats
    for name, values, _c, _d in series:
        cd.add_series(name, values)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(x),
                                Inches(y), Inches(w), Inches(h), cd)
    chart = gf.chart
    chart.has_title = False
    _style_chart(chart)
    for i, (_name, _v, color, dashed) in enumerate(series):
        s = chart.series[i]
        s.format.line.color.rgb = rgb(color)
        s.format.line.width = Pt(2.75)
        s.smooth = False
        if dashed:
            s.format.line.dash_style = 4  # dash
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    chart.legend.font.size = Pt(T.SZ_SMALL)
    chart.legend.font.name = T.FONT
    chart.legend.font.color.rgb = rgb(T.TEXT)
    va = chart.value_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = rgb(T.GREY_LIGHT)
    va.major_gridlines.format.line.width = Pt(0.75)
    va.format.line.fill.background()
    va.tick_labels.font.size = Pt(T.SZ_SMALL)
    ca = chart.category_axis
    ca.has_major_gridlines = False
    ca.format.line.color.rgb = rgb(T.GREY)
    ca.tick_labels.font.size = Pt(T.SZ_SMALL)
    return chart


# --- animation --------------------------------------------------------------
_ANIM_ID = [100]


def _nid():
    _ANIM_ID[0] += 1
    return str(_ANIM_ID[0])


def _main_seq_child(slide):
    """Return the <p:childTnLst> of the slide's main sequence, creating the
    p:timing / p:tnLst / p:seq / p:cTn skeleton on first use."""
    el = slide._element
    timing = el.find(qn("p:timing"))
    if timing is None:
        timing = el.makeelement(qn("p:timing"), {})
        el.append(timing)
    tnLst = timing.find(qn("p:tnLst"))
    if tnLst is None:
        tnLst = timing.makeelement(qn("p:tnLst"), {})
        timing.append(tnLst)
    seq = tnLst.find(qn("p:seq"))
    if seq is None:
        seq = tnLst.makeelement(qn("p:seq"),
                                {"concurrent": "1", "nextAc": "seek"})
        tnLst.append(seq)
        cTn = seq.makeelement(qn("p:cTn"),
                              {"id": "1", "dur": "indefinite",
                               "nodeType": "mainSeq"})
        seq.append(cTn)
        seq.append(seq.makeelement(qn("p:prevCondLst"), {}))
        seq.append(seq.makeelement(qn("p:nextCondLst"), {}))
        seq.find(qn("p:cTn")).append(
            seq.makeelement(qn("p:childTnLst"), {}))
    return seq.find(qn("p:cTn")).find(qn("p:childTnLst"))


# preset entry-effect filters understood by PowerPoint
FILTERS = {
    "fade": "fade",
    "zoom": "zoom(in)",
    "wipe_up": "wipe(up)",
    "wipe_left": "wipe(left)",
    "wipe_right": "wipe(right)",
    "grow": "barn(inVertical)",
    "plus": "plus",
}


def animate(slide, spid, effect="fade", delay=0, dur=500, first=False,
            trigger="after"):
    """Add an entrance animation for the shape `spid`.

    trigger: 'click'  - starts on click (used for the first item)
             'after'  - auto-chains right after the previous item
    """
    parent = _main_seq_child(slide)
    par = parent.makeelement(qn("p:par"), {})
    parent.append(par)

    cTn = par.makeelement(qn("p:cTn"),
                          {"id": _nid(), "fill": "hold",
                           "nodeType": "clickEffect" if trigger == "click"
                           else "afterEffect"})
    st = cTn.makeelement(qn("p:stCondLst"), {})
    st.append(st.makeelement(qn("p:cond"), {"delay": str(delay)}))
    cTn.append(st)
    par.append(cTn)

    kids = cTn.makeelement(qn("p:childTnLst"), {})
    cTn.append(kids)

    eff = kids.makeelement(qn("p:animEffect"),
                           {"transition": "in",
                            "filter": FILTERS.get(effect, "fade")})
    cBhvr = eff.makeelement(qn("p:cBhvr"), {})
    eff.append(cBhvr)
    c3 = cBhvr.makeelement(qn("p:cTn"), {"id": _nid(), "dur": str(dur)})
    cBhvr.append(c3)
    tgt = cBhvr.makeelement(qn("p:tgtEl"), {})
    tgt.append(tgt.makeelement(qn("p:spTgt"), {"spid": str(spid)}))
    cBhvr.append(tgt)
    kids.append(eff)
    return slide
